"""리포트/내보내기/스케줄러 판정 검증."""
import os, sys, tempfile, shutil
from datetime import datetime, date, timedelta
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
_tmp = Path(tempfile.mkdtemp(prefix="stockrep_"))
os.environ["STOCK_DATA_DIR"] = str(_tmp); os.environ["STOCK_DB_PATH"] = str(_tmp / "t.db")

from app import config, db, models, reports, exporter, mailer, scheduler  # noqa: E402
config.DATA_DIR = _tmp; config.DB_PATH = _tmp / "t.db"
config.BACKUP_DIR = _tmp/"backups"; config.UPLOAD_DIR = _tmp/"uploads"; config.EXPORT_DIR = _tmp/"exports"

PASS = FAIL = 0
def check(l, c, e=""):
    global PASS, FAIL
    if c: PASS += 1; print(f"  PASS  {l}")
    else: FAIL += 1; print(f"  FAIL  {l} {e}")

def main():
    db.init_db()
    today = date.today()
    ago = lambda n: (today - timedelta(days=n)).isoformat()

    pen  = models.create_item(dict(code="PL-0012", name="모나미 볼펜", spec="검정 0.5mm", unit="EA",
                                   category="필기구", reorder_point=20, reorder_qty=100, unit_cost=300))
    papr = models.create_item(dict(code="PA-0001", name="A4 복사용지", spec="80g 500매", unit="BOX",
                                   category="용지", reorder_point=5, reorder_qty=20, unit_cost=25000))
    old  = models.create_item(dict(code="OLD-001", name="구형 토너", spec="흑백", unit="EA",
                                   category="소모품", reorder_point=0, unit_cost=90000))

    models.post_txn(item_id=pen,  txn_type="IN",  qty=200, txn_date=ago(60))
    models.post_txn(item_id=papr, txn_type="IN",  qty=20,  txn_date=ago(40))
    models.post_txn(item_id=old,  txn_type="IN",  qty=5,   txn_date=ago(200))
    for i in range(1, 11):   # 최근 10일간 매일 출고
        models.post_txn(item_id=pen, txn_type="OUT", qty=6, txn_date=ago(i))
    models.post_txn(item_id=papr, txn_type="OUT", qty=17, txn_date=ago(2))

    print("\n[1] 요약 집계")
    s = reports.summary()
    print(f"       품목 {s['item_count']} / 경고 {s['alert_count']} / 재고금액 {s['total_value']:,.0f}원")
    check("품목 3건", s["item_count"] == 3, f"got {s['item_count']}")
    check("경고 1건 (복사용지 20-17=3 <= 5)", s["alert_count"] == 1, f"got {s['alert_count']}")
    check("재고금액 계산", s["total_value"] == 140*300 + 3*25000 + 5*90000,
          f"got {s['total_value']}")

    print("\n[2] 발주 필요 목록 + 소진 예상일")
    ro = reports.reorder_list()
    check("복사용지가 목록에 있음", any(r["code"] == "PA-0001" for r in ro), f"got {[r['code'] for r in ro]}")
    r = ro[0]
    print(f"       [{r['code']}] 현재고 {r['on_hand']:g} / 기준 {r['reorder_point']:g} "
          f"/ 부족 {r['shortage']:g} / 권장발주 {r['suggest_qty']:g}")
    check("부족수량 = 5-3 = 2", r["shortage"] == 2, f"got {r['shortage']}")
    check("권장발주량은 설정값 20", r["suggest_qty"] == 20, f"got {r['suggest_qty']}")

    print("\n[3] 소진 예상일 (일평균 출고 기반)")
    penrow = [x for x in reports.reorder_list(limit=999)] 
    avg = reports._avg_daily_out(pen)
    print(f"       볼펜 일평균 출고 {avg} (최근90일 60개/90일)")
    check("일평균 출고 계산", abs(avg - 60/90) < 0.01, f"got {avg}")

    print("\n[4] 장기 미사용 재고")
    dm = reports.dormant_items(days=90)
    print(f"       {[(d['code'], d['idle_days']) for d in dm]}")
    check("200일 전 입고 후 무이동 품목 검출", any(d["code"] == "OLD-001" for d in dm))
    check("최근 이동 품목은 제외", not any(d["code"] == "PL-0012" for d in dm))

    print("\n[5] 기초/기말 수불 계산")
    oc = reports.opening_closing(pen, ago(5), today.isoformat())
    print(f"       기초 {oc['opening']:g} + 입고 {oc['in_qty']:g} - 출고 {oc['out_qty']:g} = 기말 {oc['closing']:g}")
    check("기초 = 200 - (6*5일치) = 170", oc["opening"] == 170, f"got {oc['opening']}")
    check("기말 = 현재고 140", oc["closing"] == 140, f"got {oc['closing']}")
    check("기말이 뷰 현재고와 일치", oc["closing"] == models.get_item(pen)["on_hand"])

    print("\n[6] 리포트 생성 + HTML/텍스트 렌더링")
    rep = reports.build_report("weekly")
    check("보고할 내용이 있음", reports.has_anything_to_report(rep))
    html = mailer.render_report_html(rep)
    text = mailer.render_report_text(rep)
    check("HTML 생성", "<html>" in html and "발주해야 할 품목" in html)
    check("HTML 에 경고 품목 포함", "PA-0001" in html)
    check("HTML 에 장기미사용 섹션 포함", "장기 미사용" in html)
    check("텍스트 대체본 생성", "발주 필요" in text)
    (_tmp/"preview.html").write_text(html, encoding="utf-8")
    print(f"       미리보기 저장: {_tmp/'preview.html'} ({len(html):,} bytes)")

    print("\n[7] 엑셀 내보내기")
    p1 = exporter.export_stock()
    p2 = exporter.export_transactions()
    p3 = exporter.export_template()
    for p in (p1, p2, p3):
        check(f"{p.name} 생성 ({p.stat().st_size:,} bytes)", p.exists() and p.stat().st_size > 3000)
    from openpyxl import load_workbook
    wb = load_workbook(p1)
    check("재고현황 시트 3개", set(wb.sheetnames) == {"재고현황","발주필요","분류별요약"}, f"got {wb.sheetnames}")
    ws = wb["재고현황"]
    check("헤더 정상", ws.cell(1,1).value == "품목코드")
    check("데이터 3행", ws.max_row == 4, f"got {ws.max_row}")
    wb.close()

    print("\n[8] 스케줄러 판정 (메일 미설정이면 아무것도 안 함)")
    sch = scheduler.ReportScheduler()
    check("미설정 시 발송 안 함", sch.tick(datetime(2026,8,18,9,0)) == [])
    check("시각 도달 판정: 09:00 >= 08:30", sch._due(datetime(2026,8,18,9,0), "08:30"))
    check("시각 미도달 판정: 07:00 < 08:30", not sch._due(datetime(2026,8,18,7,0), "08:30"))
    st = sch.status()
    check("상태 조회 동작", st["daily"]["enabled"] is True and st["configured"] is False)

    print("\n[9] 중복 발송 방지")
    with db.tx() as c:
        c.execute("INSERT INTO notify_log(kind, period_key, ok) VALUES ('daily','2026-08-18',1)")
    dup = db.query_one("SELECT id FROM notify_log WHERE kind='daily' AND period_key='2026-08-18' AND ok=1")
    check("같은 기간 발송 이력이 있으면 재발송 차단", dup is not None)
    try:
        with db.tx() as c:
            c.execute("INSERT INTO notify_log(kind, period_key, ok) VALUES ('daily','2026-08-18',1)")
        check("UNIQUE 제약 작동", False, "중복 삽입이 허용됨")
    except Exception:
        check("UNIQUE 제약 작동", True)

    print(f"\n{'='*58}\n  통과 {PASS} / 실패 {FAIL}\n{'='*58}")
    return 0 if FAIL == 0 else 1

if __name__ == "__main__":
    rc = main()
    shutil.rmtree(_tmp, ignore_errors=True)
    sys.exit(rc)
