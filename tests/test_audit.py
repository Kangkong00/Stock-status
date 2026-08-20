"""
전수 점검 — 화면이 데이터를 조용히 감추지 않는지 확인한다.

지금까지 나온 사고는 대부분 같은 성격이었다. 데이터는 멀쩡한데 화면이 말없이
일부만 보여주어 사용자가 "자료가 없어졌다"고 느끼는 것이다.
그래서 이 파일은 한 가지 규칙만 본다.

    보이는 수 == 실제 수,  아니면 반드시 몇 건 중 몇 건인지 알린다.
"""
import json, os, re, shutil, sys, tempfile
from datetime import date, timedelta
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
_tmp = Path(tempfile.mkdtemp(prefix="stockaudit_"))
os.environ["STOCK_DATA_DIR"] = str(_tmp)
os.environ["STOCK_DB_PATH"] = str(_tmp / "t.db")

from app import config, db, exporter, models, reports  # noqa: E402
config.DATA_DIR = _tmp
config.DB_PATH = _tmp / "t.db"
config.BACKUP_DIR = _tmp / "b"
config.UPLOAD_DIR = _tmp / "u"
config.EXPORT_DIR = _tmp / "e"
from app.web import create_app  # noqa: E402

PASS = FAIL = 0


def check(label, cond, extra=""):
    global PASS, FAIL
    if cond:
        PASS += 1
        print(f"  PASS  {label}")
    else:
        FAIL += 1
        print(f"  FAIL  {label}  {extra}")


DATE_TD = re.compile(r'<td class="muted small nowrap">2\d{3}-\d\d-\d\d</td>')
HEAD_TOTAL = re.compile(r"<b>([\d,]+)건</b>")


def shown_rows(html: str) -> int:
    return len(DATE_TD.findall(html))


def head_total(html: str) -> int | None:
    m = HEAD_TOTAL.search(html)
    return int(m.group(1).replace(",", "")) if m else None


def seed() -> tuple[int, int]:
    """3년치 이력을 가진 품목을 만든다. 쪽 나눔이 필요한 규모여야 의미가 있다."""
    db.init_db()
    soda = models.create_item(dict(code="CH-0001", name="가성소다 고상(98%)", unit="KG",
                                   category="약품류", pack_unit="포", pack_size=25,
                                   reorder_point=100))
    pen = models.create_item(dict(code="PC-0001", name="볼펜", unit="PC",
                                  category="사무용품", reorder_point=10))
    base = date(2023, 1, 1)
    for i in range(600):                       # 200건씩 3쪽이 나오는 규모
        models.post_txn(item_id=soda, txn_type="IN" if i % 2 == 0 else "OUT", qty=25,
                        txn_date=(base + timedelta(days=i * 2)).isoformat(),
                        partner="한국문구", doc_no=f"A-{i:04d}")
    for i in range(30):
        models.post_txn(item_id=pen, txn_type="IN", qty=5,
                        txn_date=(base + timedelta(days=i)).isoformat())
    for i in range(40):                        # 검색 상한 확인용
        models.create_item(dict(code=f"VL-{i:04d}", name=f"밸브 {i}호",
                                unit="PC", category="밸브류"))
    return soda, pen


def main() -> int:
    soda, pen = seed()
    app = create_app()
    app.config["TESTING"] = True
    c = app.test_client()

    TOTAL = models.txn_count()
    SODA = models.txn_count(item_id=soda)
    print(f"\n[준비] 품목 {db.scalar('SELECT COUNT(*) FROM items')}건 / 전표 {TOTAL}건 "
          f"(가성소다 {SODA}건, 2023~2026)")

    print("\n[1] 입출고 내역이 과거를 감추지 않는다")
    h = c.get("/history").data.decode()
    check(f"머리말이 전체 건수를 밝힌다 ({head_total(h)})", head_total(h) == TOTAL,
          f"got {head_total(h)} / 실제 {TOTAL}")
    check("기본 기간이 '전체' 다", shown_rows(h) >= 200, f"{shown_rows(h)}행")
    check("쪽 넘기기가 있다", "다음 ›" in h)
    # "전체 2,984건 중 <b>1~200번째</b>를 보고 있습니다" 처럼 태그로 끊겨 있다
    check("몇 번째를 보고 있는지 알려준다",
          re.search(r"전체[^<]*건 중\s*<b>[\d,]+~[\d,]+번째</b>를 보고 있습니다", h) is not None,
          "쪽 위치 안내 문구 없음")

    print("\n[2] ★ 검색이 전 기간을 훑는다")
    h = c.get("/history?keyword=가성소다").data.decode()
    check(f"'가성소다' → {head_total(h)}건 (실제 {SODA})", head_total(h) == SODA,
          f"got {head_total(h)}")
    for kw, label in [("A-0003", "전표번호"), ("한국문구", "거래처"), ("CH-0001", "품목코드")]:
        hh = c.get(f"/history?keyword={kw}").data.decode()
        check(f"{label}로도 찾는다 ({head_total(hh)}건)", (head_total(hh) or 0) >= 1)

    print("\n[3] 쪽을 다 넘기면 전부 볼 수 있다")
    seen, page = 0, 1
    while page <= 40:
        hh = c.get(f"/history?keyword=가성소다&page={page}").data.decode()
        n = shown_rows(hh)
        seen += n
        m = re.search(r"<b>(\d+)</b> / (\d+)", hh)
        if not m or int(m.group(1)) >= int(m.group(2)) or n == 0:
            break
        page += 1
    check(f"{page}쪽을 넘겨 {seen}건 전부 확인 (실제 {SODA})", seen == SODA, f"got {seen}")

    print("\n[4] 기간 버튼")
    counts = {}
    for p, label in [("all", "전체"), ("365", "1년"), ("90", "90일"),
                     ("30", "30일"), ("7", "7일"), ("today", "오늘")]:
        r = c.get(f"/history?period={p}")
        counts[p] = head_total(r.data.decode())
        check(f"{label} 버튼 → {counts[p]}건", r.status_code == 200, f"HTTP {r.status_code}")
    check("기간이 좁을수록 건수가 줄어든다",
          counts["all"] >= counts["365"] >= counts["90"] >= counts["30"] >= counts["7"],
          str(counts))

    print("\n[5] 품목 상세가 이력을 감추면 알린다")
    h = c.get(f"/stock/{soda}").data.decode()
    check(f"전체 건수를 밝힌다 ({SODA}건)", f"{SODA:,}건" in h or f">{SODA}건" in h)
    if shown_rows(h) < SODA:
        check("잘렸으면 '나머지 N건 보기' 를 준다", "나머지" in h and "보기" in h,
              f"{shown_rows(h)}/{SODA} 인데 안내 없음")
    check("전체 내역 링크가 그 품목으로 걸린다", f"item_id={soda}" in h)
    h2 = c.get(f"/history?item_id={soda}&voided=1").data.decode()
    check(f"그 링크로 가면 {SODA}건 전부 나온다", head_total(h2) == SODA, f"got {head_total(h2)}")

    print("\n[6] 품목 검색이 잘리면 알린다")
    d = json.loads(c.get("/api/search?q=밸브").data)
    total = models.stock_count(keyword="밸브")
    check(f"전체 건수를 함께 준다 ({d.get('total')} / 실제 {total})", d.get("total") == total)
    check("예전 15건 상한보다 많이 준다", len(d["items"]) >= 20, f"{len(d['items'])}건")
    if len(d["items"]) < total:
        check("잘렸다고 표시한다", d.get("truncated") is True)

    print("\n[7] 등록 화면의 입출고 내역 패널")
    h = c.get("/register").data.decode()
    check("저장 버튼 아래에 조회 패널이 있다", "hTable" in h and "histLoad" in h)
    check("기간 버튼이 있다", 'data-hp="all"' in h and 'data-hp="30"' in h)
    check("전체 화면으로 가는 링크가 있다", "전체 화면으로" in h)
    d = json.loads(c.get("/api/history?limit=5").data)
    check(f"조회 API 가 전체 건수를 준다 ({d['total']} / 실제 {TOTAL})", d["total"] == TOTAL)
    check("더 있으면 알린다", d["has_more"] is True)
    d = json.loads(c.get(f"/api/history?item_id={soda}&limit=5").data)
    check(f"품목으로 좁히면 그 품목만 ({d['total']} / 실제 {SODA})", d["total"] == SODA)
    d2 = json.loads(c.get(f"/api/history?item_id={soda}&type=OUT&limit=5").data)
    out_n = models.txn_count(item_id=soda, txn_type="OUT", include_voided=True)
    check(f"구분 거르기가 동작한다 (출고 {d2['total']} / 실제 {out_n})", d2["total"] == out_n)
    d3 = json.loads(c.get(f"/api/history?item_id={soda}&period=30&limit=5").data)
    check(f"기간 거르기가 동작한다 (30일 {d3['total']} < 전체 {SODA})", d3["total"] < SODA)
    d4 = json.loads(c.get(f"/api/history?item_id={soda}&limit=5&offset=5").data)
    check("이어서 더 가져오기가 동작한다",
          d4["rows"] and d4["rows"][0]["id"] != d["rows"][0]["id"])

    print("\n[8] 대시보드가 실제 건수를 말한다")
    h = c.get("/").data.decode()
    alert = reports.summary()["alert_count"]
    check(f"발주 필요 {alert}건을 그대로 표시", f"{alert}건" in h or f">{alert}<" in h)
    if alert > 12:
        check("12건만 보여줄 때 나머지를 알린다", "나머지" in h and "보기" in h)

    print("\n[9] 재고현황은 전 품목을 보여준다")
    h = c.get("/stock").data.decode()
    shown = len(re.findall(r'href="/stock/\d+"', h))
    active = db.scalar("SELECT COUNT(*) FROM items WHERE active=1")
    check(f"활성 {active}건 전부 표시 (화면 {shown})", shown >= active, f"{shown}/{active}")

    print("\n[10] 엑셀에 전 건이 담긴다")
    from openpyxl import load_workbook
    p = exporter.export_transactions()
    wb = load_workbook(p)
    n = wb["입출고내역"].max_row - 1
    wb.close()
    allrows = db.scalar("SELECT COUNT(*) FROM transactions")
    check(f"입출고내역 엑셀 {n}행 == 전표 {allrows}건", n == allrows)
    p = exporter.export_stock()
    wb = load_workbook(p)
    n = wb["재고현황"].max_row - 1
    wb.close()
    check(f"재고현황 엑셀 {n}행 == 활성 {active}건", n == active)

    print("\n[11] 데이터 정합성")
    mis = sum(
        1 for row in models.list_stock(include_inactive=True)
        if abs(db.scalar("SELECT COALESCE(SUM(signed_qty),0) FROM transactions "
                         "WHERE item_id=? AND voided=0", (row["item_id"],), 0)
               - row["on_hand"]) > 1e-9)
    check(f"원장 합산 == 현재고 (불일치 {mis})", mis == 0)
    check("고아 전표 없음", db.scalar(
        "SELECT COUNT(*) FROM transactions WHERE item_id NOT IN (SELECT id FROM items)", (), 0) == 0)
    check("일자 없는 전표 없음", db.scalar(
        "SELECT COUNT(*) FROM transactions WHERE txn_date IS NULL OR txn_date=''", (), 0) == 0)

    print("\n[12] 모든 화면이 열린다")
    for p in ["/", "/stock", "/stock?status=ALERT", "/stock?keyword=밸브", "/stock?all=1",
              "/register", "/pending", "/history", "/history?period=30", "/stocktake",
              "/upload", "/settings", "/report/preview", "/report/preview?kind=weekly",
              "/api/health", "/api/categories", "/api/next-code?category=약품류",
              "/api/history", f"/stock/{soda}", "/export/stock", "/export/template"]:
        r = c.get(p)
        check(f"{p} → {r.status_code}", r.status_code == 200)

    print(f"\n{'=' * 60}\n  통과 {PASS} / 실패 {FAIL}\n{'=' * 60}")
    return 0 if FAIL == 0 else 1


if __name__ == "__main__":
    rc = main()
    shutil.rmtree(_tmp, ignore_errors=True)
    sys.exit(rc)
