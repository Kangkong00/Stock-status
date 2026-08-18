"""엑셀 내보내기. 관리는 시스템이, 보고서는 엑셀로."""
from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import config, db, models, reports

HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=10)
TITLE_FONT = Font(bold=True, size=14)
LOW_FILL = PatternFill("solid", fgColor="FFF3CD")
OUT_FILL = PatternFill("solid", fgColor="F8D7DA")
THIN = Side(style="thin", color="D0D7DE")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _sheet(wb: Workbook, title: str, headers: list[str], first: bool = False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title[:31]
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER
    ws.freeze_panes = "A2"
    return ws


def _autosize(ws, min_w: int = 8, max_w: int = 42) -> None:
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        width = max((len(str(c.value or "")) for c in col), default=0)
        # 한글은 폭을 더 먹는다
        korean = max((sum(1 for ch in str(c.value or "") if ord(ch) > 0x1100) for c in col), default=0)
        ws.column_dimensions[letter].width = min(max(width + korean * 0.7 + 3, min_w), max_w)


def _stamp(name: str) -> Path:
    config.ensure_dirs()
    return config.EXPORT_DIR / f"{name}_{datetime.now():%Y%m%d_%H%M%S}.xlsx"


STATUS_KO = {"OK": "정상", "LOW": "재고기준 도달", "OUT": "품절"}


def export_stock(path: Path | None = None, *, include_inactive: bool = False) -> Path:
    """재고현황표. 기존 엑셀 보고서를 대체하는 기본 산출물."""
    path = path or _stamp("재고현황")
    wb = Workbook()

    ws = _sheet(wb, "재고현황", [
        "품목코드", "품목명", "규격", "단위", "분류", "보관위치",
        "현재고", "재고기준", "과부족", "상태", "단가", "재고금액",
        "누적입고", "누적출고", "최종입고일", "최종출고일",
    ], first=True)

    for r in models.list_stock(include_inactive=include_inactive):
        shortage = (r["on_hand"] or 0) - (r["reorder_point"] or 0)
        ws.append([
            r["code"], r["name"], r["spec"], r["unit"], r["category"], r["location"],
            r["on_hand"], r["reorder_point"], shortage, STATUS_KO.get(r["stock_status"], r["stock_status"]),
            r["unit_cost"], r["stock_value"], r["total_in"], r["total_out"],
            r["last_in_date"] or "", r["last_out_date"] or "",
        ])
        if r["stock_status"] == "OUT":
            for c in ws[ws.max_row]:
                c.fill = OUT_FILL
        elif r["stock_status"] == "LOW":
            for c in ws[ws.max_row]:
                c.fill = LOW_FILL

    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.border = BORDER
        for idx in (7, 8, 9, 11, 12, 13, 14):
            row[idx - 1].number_format = "#,##0.##"
    _autosize(ws)

    # 발주 필요 목록
    ws2 = _sheet(wb, "발주필요", [
        "품목코드", "품목명", "규격", "단위", "현재고", "재고기준",
        "부족수량", "권장발주량", "예상금액", "일평균출고", "소진예상(일)",
    ])
    for r in reports.reorder_list():
        ws2.append([r["code"], r["name"], r["spec"], r["unit"], r["on_hand"],
                    r["reorder_point"], r["shortage"], r["suggest_qty"],
                    r["suggest_amount"], r["avg_daily_out"],
                    r["days_left"] if r["days_left"] is not None else ""])
    _autosize(ws2)

    # 분류별 요약
    ws3 = _sheet(wb, "분류별요약", ["분류", "품목수", "총수량", "재고금액", "경고품목수"])
    for r in reports.category_breakdown():
        ws3.append([r["category"], r["item_count"], r["qty"], r["value"], r["alert_count"]])
    _autosize(ws3)

    wb.save(path)
    return path


def export_transactions(date_from: str = "", date_to: str = "", path: Path | None = None) -> Path:
    """입출고 내역서 (수불부)."""
    path = path or _stamp("입출고내역")
    wb = Workbook()
    ws = _sheet(wb, "입출고내역", [
        "일자", "구분", "품목코드", "품목명", "규격", "단위", "수량",
        "단가", "금액", "거래처", "전표번호", "원문상품명", "비고", "취소",
    ], first=True)

    kind = {"IN": "입고", "OUT": "출고", "ADJ": "조정"}
    rows = models.txn_history(date_from=date_from, date_to=date_to,
                              include_voided=True, limit=100000)
    for t in rows:
        ws.append([
            t["txn_date"], kind.get(t["txn_type"], t["txn_type"]), t["code"], t["name"],
            t["spec"], t["unit"], t["signed_qty"], t["unit_cost"],
            abs(t["signed_qty"]) * (t["unit_cost"] or 0), t["partner"], t["doc_no"],
            t["raw_name"], t["memo"], "취소" if t["voided"] else "",
        ])
        if t["voided"]:
            for c in ws[ws.max_row]:
                c.font = Font(strike=True, color="999999")

    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.border = BORDER
        for idx in (7, 8, 9):
            row[idx - 1].number_format = "#,##0.##"
    _autosize(ws)
    wb.save(path)
    return path


def export_stocktake(stocktake_id: int, path: Path | None = None) -> Path:
    """실사 카운트 시트. 인쇄해서 창고에서 손으로 적고 다시 올리는 용도."""
    st = db.query_one("SELECT * FROM stocktakes WHERE id = ?", (stocktake_id,))
    if st is None:
        raise models.DomainError("실사 건을 찾을 수 없습니다.")
    path = path or _stamp(f"재고실사_{st['take_date']}")

    wb = Workbook()
    ws = _sheet(wb, "재고실사", [
        "품목코드", "품목명", "규격", "단위", "보관위치", "전산재고", "실사수량", "차이", "비고",
    ], first=True)
    for r in models.stocktake_lines(stocktake_id):
        counted = r["counted_qty"]
        ws.append([r["code"], r["name"], r["spec"], r["unit"], r["location"],
                   r["system_qty"], counted if counted is not None else "",
                   (counted - r["system_qty"]) if counted is not None else "", r["memo"]])
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.border = BORDER
    _autosize(ws)
    wb.save(path)
    return path


def export_template(path: Path | None = None) -> Path:
    """빈 입력 양식. 이 양식대로 채워서 올리면 100% 자동 인식된다."""
    path = path or config.EXPORT_DIR / "업로드_양식.xlsx"
    config.ensure_dirs()
    wb = Workbook()

    ws = _sheet(wb, "품목등록", [
        "품목코드", "품목명", "규격", "단위", "분류", "바코드",
        "보관위치", "재고기준", "권장발주량", "매입단가", "현재고", "비고",
    ], first=True)
    ws.append(["PL-0012", "모나미 볼펜", "검정 0.5mm", "EA", "필기구", "",
               "A-1", 20, 100, 300, 145, "예시 행입니다. 지우고 쓰세요."])
    _autosize(ws)

    ws2 = _sheet(wb, "입출고등록", [
        "일자", "품목코드", "상품명", "입고", "출고", "단가", "거래처", "전표번호", "비고",
    ])
    ws2.append(["2026-08-18", "PL-0012", "모나미 볼펜 검정 0.5mm", 100, "", 300,
                "한국문구", "A-1001", "예시 행입니다. 지우고 쓰세요."])
    ws2.append(["2026-08-18", "", "A4 복사용지 80g 500매", "", 5, "", "", "O-2001",
                "품목코드를 몰라도 상품명만 있으면 자동 매칭됩니다"])
    _autosize(ws2)

    ws3 = wb.create_sheet("사용법")
    guide = [
        ["업로드 양식 사용법"],
        [],
        ["1. [품목등록] 시트"],
        ["   · 품목코드와 품목명은 필수입니다."],
        ["   · '재고기준'에 적은 수량 이하로 떨어지면 대시보드와 메일로 경고가 옵니다."],
        ["   · '현재고' 열을 채우면 기초재고가 입고 전표로 자동 생성됩니다."],
        ["   · 이미 있는 품목코드는 내용이 갱신됩니다(덮어쓰기)."],
        [],
        ["2. [입출고등록] 시트"],
        ["   · 품목코드를 적으면 즉시 매칭됩니다."],
        ["   · 품목코드 없이 상품명만 적어도 됩니다. 시스템이 이름으로 찾아냅니다."],
        ["   · 찾지 못한 이름은 '확인 대기' 목록에 쌓입니다. 화면에서 클릭 한 번으로 연결하면"],
        ["     그 이름을 기억해서 다음부터는 자동으로 처리합니다."],
        ["   · 입고와 출고는 각각의 열에 적습니다. 한 행에 둘 다 적어도 됩니다."],
        [],
        ["3. 열 이름을 바꿔도 됩니다"],
        ["   · 업로드 화면에서 어느 열이 무엇인지 직접 지정할 수 있습니다."],
        ["   · 반영 전에 '시험 실행'으로 결과를 미리 확인하세요."],
    ]
    for line in guide:
        ws3.append(line)
    ws3["A1"].font = TITLE_FONT
    ws3.column_dimensions["A"].width = 90

    wb.save(path)
    return path
